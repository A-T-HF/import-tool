"""
Mews XLSX Reader

Supports two formats:
1. Mews Reservierungsbericht export  (sheets: Reservierungen, Parameter, Nächte)
2. Mews Import Template              (sheets: Reservations, Customers, Companies)

Usage:
    from mews_reader import is_mews_xlsx, read_mews
    if is_mews_xlsx(path):
        rows = read_mews(path, "guest")
"""

import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

from guests_importer import is_company_name, derive_name_from_email


# ── Format signatures ─────────────────────────────────────────────────────────

_MEWS_SHEETS         = {"Reservierungen", "Parameter", "Nächte"}
_MEWS_COLS           = {"Nummer", "Anreise", "Abreise", "Kennung", "Raumkategorie"}
_MEWS_TEMPLATE_SHEETS = {"Reservations", "Customers", "Companies"}

TEMPLATE_HEADER_ROW = 4   # first 3 rows: title / hint / blank


# ── Detection ─────────────────────────────────────────────────────────────────

def is_mews_xlsx(file_or_path) -> bool:
    """Return True if the file is a Mews XLSX (export report or import template)."""
    try:
        _seek(file_or_path)
        wb = openpyxl.load_workbook(file_or_path, read_only=True, data_only=True)
        sheets = set(wb.sheetnames)
        if _MEWS_SHEETS.issubset(sheets):
            ws = wb["Reservierungen"]
            header_row = next(ws.iter_rows(max_row=1, values_only=True), None)
            if header_row and _MEWS_COLS.issubset(set(header_row)):
                return True
        if _MEWS_TEMPLATE_SHEETS.issubset(sheets):
            return True
        return False
    except Exception:
        return False


def _is_template(wb) -> bool:
    return _MEWS_TEMPLATE_SHEETS.issubset(set(wb.sheetnames))


def _seek(file_or_path):
    if hasattr(file_or_path, "seek"):
        file_or_path.seek(0)


# ── German country names → ISO 3166-1 alpha-2 ────────────────────────────────

_DE_COUNTRY = {
    "deutschland": "DE",
    "österreich": "AT",
    "schweiz": "CH",
    "frankreich": "FR",
    "italien": "IT",
    "spanien": "ES",
    "portugal": "PT",
    "niederlande": "NL",
    "belgien": "BE",
    "luxemburg": "LU",
    "dänemark": "DK",
    "schweden": "SE",
    "norwegen": "NO",
    "finnland": "FI",
    "island": "IS",
    "irland": "IE",
    "vereinigtes königreich großbritannien und nordirland": "GB",
    "vereinigtes königreich": "GB",
    "großbritannien": "GB",
    "vereinigte staaten von amerika": "US",
    "vereinigte staaten": "US",
    "usa": "US",
    "kanada": "CA",
    "australien": "AU",
    "neuseeland": "NZ",
    "japan": "JP",
    "china": "CN",
    "südkorea": "KR",
    "indien": "IN",
    "brasilien": "BR",
    "argentinien": "AR",
    "mexiko": "MX",
    "türkei": "TR",
    "russland": "RU",
    "ukraine": "UA",
    "polen": "PL",
    "tschechien": "CZ",
    "tschechische republik": "CZ",
    "slowakei": "SK",
    "ungarn": "HU",
    "rumänien": "RO",
    "bulgarien": "BG",
    "kroatien": "HR",
    "serbien": "RS",
    "slowenien": "SI",
    "bosnien und herzegowina": "BA",
    "griechenland": "GR",
    "zypern": "CY",
    "malta": "MT",
    "lettland": "LV",
    "litauen": "LT",
    "estland": "EE",
    "weißrussland": "BY",
    "moldawien": "MD",
    "georgien": "GE",
    "armenien": "AM",
    "aserbaidschan": "AZ",
    "israel": "IL",
    "saudi-arabien": "SA",
    "vereinigte arabische emirate": "AE",
    "ägypten": "EG",
    "marokko": "MA",
    "tunesien": "TN",
    "südafrika": "ZA",
    "nigeria": "NG",
    "kenia": "KE",
    "singapur": "SG",
    "thailand": "TH",
    "indonesien": "ID",
    "malaysia": "MY",
    "philippinen": "PH",
    "vietnam": "VN",
    "kolumbien": "CO",
    "chile": "CL",
    "peru": "PE",
    "venezuela": "VE",
    "liechtenstein": "LI",
    "monaco": "MC",
    "andorra": "AD",
    "san marino": "SM",
    "nordmazedonien": "MK",
    "albanien": "AL",
    "montenegro": "ME",
}


def _map_nationality(german_name) -> str | None:
    if not german_name:
        return None
    return _DE_COUNTRY.get(str(german_name).strip().lower())


# ── Helpers ───────────────────────────────────────────────────────────────────

def _clean(v) -> str:
    return (str(v) if v is not None else "").strip()


def _fmt_date(v) -> str | None:
    """Convert openpyxl datetime/date, DD/MM/YYYY string, or YYYY-MM-DD to YYYY-MM-DD."""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    s = _clean(v)
    if not s:
        return None
    # DD/MM/YYYY  (Mews template export format)
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", s)
    if m:
        return f"{m.group(3)}-{m.group(2):0>2}-{m.group(1):0>2}"
    return s


def _parse_person_name(nachname, vorname) -> tuple[str, str]:
    """
    Returns (first_name, last_name).

    Mews writes 'Company, Person Name' in Nachname when booked via company.
    In that case Vorname contains 'First Last' for the real person.
    """
    n = _clean(nachname)
    v = _clean(vorname)
    if "," in n and v:
        parts = v.rsplit(" ", 1)
        if len(parts) == 2:
            return parts[0].strip(), parts[1].strip()
        return "", v
    return v, n


def _group_display_name(gruppenname) -> str:
    """Strip Mews suffix '-NN-NN-XXXX' from group/booking name."""
    return re.sub(r"-\d+-\d+-[A-F0-9]{4}$", "", _clean(gruppenname)).strip()


def _norm_header(h) -> str:
    """Normalize column header: strip trailing ' *' and whitespace."""
    return re.sub(r"\s*\*+\s*$", "", str(h) if h is not None else "").strip()


# ── Reservierungsbericht (export) sheet loading ───────────────────────────────

def _load_rows_from_wb(wb) -> list[dict]:
    """Load 'Reservierungen' sheet from an already-opened workbook."""
    ws = wb["Reservierungen"]
    raw = list(ws.iter_rows(values_only=True))
    if not raw:
        return []
    headers = [str(h) if h is not None else "" for h in raw[0]]
    anreise_idx = headers.index("Anreise") if "Anreise" in headers else -1
    rows = []
    for row in raw[1:]:
        first = str(row[0]).strip() if row[0] is not None else ""
        if not first.isdigit():
            continue
        if anreise_idx >= 0:
            anreise_val = row[anreise_idx]
            if not isinstance(anreise_val, (date, datetime)):
                continue
        rows.append(dict(zip(headers, row)))
    return rows


def _load_rows(file_or_path) -> list[dict]:
    wb = openpyxl.load_workbook(file_or_path, read_only=True, data_only=True)
    return _load_rows_from_wb(wb)


# ── Import Template sheet loading ─────────────────────────────────────────────

def _load_template_rows(wb, sheet_name: str) -> list[dict]:
    """Load a template sheet (header at row TEMPLATE_HEADER_ROW, data below)."""
    ws = wb[sheet_name]
    raw = list(ws.iter_rows(values_only=True))
    if len(raw) < TEMPLATE_HEADER_ROW:
        return []
    headers = [_norm_header(h) for h in raw[TEMPLATE_HEADER_ROW - 1]]
    rows = []
    for row in raw[TEMPLATE_HEADER_ROW:]:
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        rows.append(dict(zip(headers, row)))
    return rows


# ── Reservierungsbericht guest/company/reservation readers ────────────────────

def _read_guests(rows: list[dict]) -> list[dict]:
    seen: dict[tuple, bool] = {}
    result = []
    for row in rows:
        first, last = _parse_person_name(row.get("Nachname"), row.get("Vorname"))
        first, last = first.strip(), last.strip()

        if last and is_company_name(last):
            continue

        email = _clean(row.get("E-Mail"))

        name_source = "original" if first else "none"
        if not first and last and email:
            derived = derive_name_from_email(email, last)
            if derived:
                first = derived
                name_source = "derived_from_email"

        key = (first.lower(), last.lower(), email.lower())
        if key in seen:
            continue
        seen[key] = True

        sys: dict = {}
        if name_source == "derived_from_email":
            sys["first_name"] = ""
        elif not first:
            sys["first_name"] = ""
        if not last:
            sys["last_name"] = ""

        nat_raw = _clean(row.get("Staatsangehörigkeit des Gastes"))
        nat = _map_nationality(nat_raw)
        if nat:
            sys["nationality"] = nat_raw

        g: dict = {
            "first_name":      first or "-",
            "last_name":       last  or "-",
            "_name_source":    name_source,
            "_system_changes": sys,
        }
        if email:
            g["email"] = email
        phone = _clean(row.get("Telefon"))
        if phone:
            g["phone"] = phone
        if nat:
            g["nationality"] = nat
        mews_id = _clean(row.get("Kennung"))
        if mews_id:
            g["_mews_id"] = mews_id
        result.append(g)
    return result


def _read_companies(rows: list[dict]) -> list[dict]:
    companies: dict[str, dict] = {}

    for row in rows:
        firma = _clean(row.get("Firma"))
        if firma and firma not in companies:
            companies[firma] = {
                "name":     firma,
                "_mews_id": _clean(row.get("Firma Kennung")) or None,
            }

    group_rows: dict[str, list[dict]] = {}
    for row in rows:
        gname = _clean(row.get("Gruppenname"))
        if gname:
            group_rows.setdefault(gname, []).append(row)

    for gname, glist in group_rows.items():
        if len(glist) < 2:
            continue
        display = _group_display_name(gname)
        if display and display not in companies:
            sys: dict = {}
            if display != gname:
                sys["name"] = gname
            companies[display] = {
                "name":            display,
                "_mews_id":        _clean(glist[0].get("Kennung")) or None,
                "_system_changes": sys,
            }

    return [{k: v for k, v in c.items() if v} for c in companies.values()]


def _read_reservations(rows: list[dict]) -> list[dict]:
    result = []
    for row in rows:
        raw_nachname = _clean(row.get("Nachname"))
        raw_vorname  = _clean(row.get("Vorname"))
        first, last  = _parse_person_name(raw_nachname, raw_vorname)
        first, last  = first.strip(), last.strip()

        email = _clean(row.get("E-Mail"))

        if not first and last and email:
            derived = derive_name_from_email(email, last)
            if derived:
                first = derived

        sys: dict = {}
        if "," in raw_nachname and raw_vorname:
            sys["first_name"] = raw_nachname
        elif not first:
            sys["first_name"] = ""
        if not last:
            sys["last_name"] = ""

        r: dict = {
            "first_name":      first or "-",
            "last_name":       last  or "-",
            "Check In":        _fmt_date(row.get("Anreise")),
            "Check Out":       _fmt_date(row.get("Abreise")),
            "Zimmer":          _clean(row.get("Raumnummer")),
            "Zimmertyp":       _clean(row.get("Raumkategorie")),
            "_system_changes": sys,
        }
        if email:
            r["email"] = email
        summe = row.get("Gesamtbetrag")
        if summe is not None and summe != "":
            r["Summe"] = str(summe)
        status = _clean(row.get("Status"))
        if status:
            r["Status"] = status
        mews_id = _clean(row.get("Kennung"))
        if mews_id:
            r["_mews_id"] = mews_id
        result.append({k: v for k, v in r.items() if v is not None and v != ""})
    return result


# ── Import Template guest/company/reservation readers ─────────────────────────

def _read_template_guests(rows: list[dict]) -> list[dict]:
    """Customers sheet → HotelFriend guest rows."""
    seen: dict[tuple, bool] = {}
    result = []
    for row in rows:
        first = _clean(row.get("First Name"))
        last  = _clean(row.get("Last Name"))
        if not last:
            continue
        email = _clean(row.get("Email"))
        key = (first.lower(), last.lower(), email.lower())
        if key in seen:
            continue
        seen[key] = True

        g: dict = {
            "first_name": first or "-",
            "last_name":  last,
        }
        if email:
            g["email"] = email
        phone = _clean(row.get("Phone"))
        if phone:
            g["phone"] = phone
        nat = _clean(row.get("Nationality Code"))
        if nat:
            g["nationality"] = nat
        bd = _fmt_date(row.get("Birth Date"))
        if bd:
            g["birthday"] = bd
        addr = _clean(row.get("Address Line 1"))
        if addr:
            g["address"] = addr
        city = _clean(row.get("City"))
        if city:
            g["city"] = city
        zip_ = _clean(row.get("ZIP"))
        if zip_:
            g["zip_code"] = zip_
        country = _clean(row.get("Country Code"))
        if country:
            g["country"] = country
        result.append(g)
    return result


def _read_template_companies(rows: list[dict]) -> list[dict]:
    """Companies sheet → HotelFriend company rows."""
    result = []
    for row in rows:
        name = _clean(row.get("Name"))
        if not name:
            continue
        c: dict = {"name": name}
        addr = _clean(row.get("Address Line 1"))
        if addr:
            c["address"] = addr
        city = _clean(row.get("City"))
        if city:
            c["city"] = city
        zip_ = _clean(row.get("ZIP"))
        if zip_:
            c["zip_code"] = zip_
        country = _clean(row.get("Country Code"))
        if country:
            c["country"] = country
        email = _clean(row.get("Contact Email"))
        if email:
            c["email"] = email
        phone = _clean(row.get("Contact Phone"))
        if phone:
            c["phone"] = phone
        tax = _clean(row.get("Tax Identifier"))
        if tax:
            c["tax_id"] = tax
        result.append(c)
    return result


def _read_template_reservations(rows: list[dict]) -> list[dict]:
    """Reservations sheet → HotelFriend reservation rows."""
    result = []
    for row in rows:
        first = _clean(row.get("First Name"))
        last  = _clean(row.get("Last Name"))
        email = _clean(row.get("Email"))

        r: dict = {
            "first_name": first or "-",
            "last_name":  last  or "-",
            "Check In":   _fmt_date(row.get("Reservation Start")),
            "Check Out":  _fmt_date(row.get("Reservation End")),
            "Zimmertyp":  _clean(row.get("Resource Category")),
            "Zimmer":     _clean(row.get("Resource")),
        }
        if email:
            r["email"] = email
        summe = row.get("Total Price")
        if summe is not None and str(summe).strip():
            r["Summe"] = str(summe)
        currency = _clean(row.get("Currency Code"))
        if currency:
            r["Währung"] = currency
        company = _clean(row.get("Company Name"))
        if company:
            r["Firma"] = company
        conf = _clean(row.get("Confirmation Code"))
        if conf:
            r["_confirmation_code"] = conf
        result.append({k: v for k, v in r.items() if v is not None and v != ""})
    return result


# ── Entry point ───────────────────────────────────────────────────────────────

def read_mews(file_or_path, entity_type: str) -> list[dict]:
    """
    Read a Mews XLSX and return rows pre-mapped to HotelFriend import format.

    Supports:
    - Mews Reservierungsbericht export (Reservierungen / Parameter / Nächte)
    - Mews Import Template             (Reservations / Customers / Companies)

    file_or_path: path string, Path, or file-like object (BytesIO)
    entity_type:  "guest" | "company" | "reservation"
    """
    _seek(file_or_path)
    wb = openpyxl.load_workbook(file_or_path, read_only=True, data_only=True)

    if _is_template(wb):
        if entity_type == "guest":
            return _read_template_guests(_load_template_rows(wb, "Customers"))
        if entity_type == "company":
            return _read_template_companies(_load_template_rows(wb, "Companies"))
        if entity_type == "reservation":
            return _read_template_reservations(_load_template_rows(wb, "Reservations"))
    else:
        rows = _load_rows_from_wb(wb)
        if entity_type == "guest":
            return _read_guests(rows)
        if entity_type == "company":
            return _read_companies(rows)
        if entity_type == "reservation":
            return _read_reservations(rows)

    raise ValueError(f"Unbekannter Entitätstyp: {entity_type}")
