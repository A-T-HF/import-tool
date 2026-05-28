"""
BookingList XLSX Reader

Supports the "BookingList-YYYY-MM-DD.xlsx" format exported by booking tools
like Lodgify, Smoobu, and similar portal aggregators. Characteristic columns:
Position, Gast, Anreise, Abreise, Unterkunft.

The "Gast" column contains "Vorname Nachname" as a single string — this reader
splits it automatically.

Usage:
    from bookinglist_reader import is_bookinglist_xlsx, read_bookinglist
    if is_bookinglist_xlsx(path):
        rows = read_bookinglist(path, "reservation")
"""

import re
from datetime import datetime

import openpyxl

from guests_importer import is_company_name
from transformer import _STATUS_MAP, transform_country


# ── Format signature ──────────────────────────────────────────────────────────

_SIGNATURE_COLS = {"Position", "Gast", "Anreise", "Abreise", "Unterkunft"}


# ── Detection ─────────────────────────────────────────────────────────────────

def is_bookinglist_xlsx(file_or_path) -> bool:
    """Return True if the file looks like a BookingList XLSX export."""
    try:
        _seek(file_or_path)
        wb = openpyxl.load_workbook(file_or_path, read_only=True, data_only=True)
        ws = wb.active
        header_row = next(ws.iter_rows(max_row=1, values_only=True), None)
        if header_row:
            headers = {str(h) for h in header_row if h is not None}
            return _SIGNATURE_COLS.issubset(headers)
        return False
    except Exception:
        return False


def _seek(file_or_path):
    if hasattr(file_or_path, "seek"):
        file_or_path.seek(0)


# ── Helpers ───────────────────────────────────────────────────────────────────

def _clean(v) -> str:
    return (str(v) if v is not None else "").strip()


_BL_DATE_FMTS = ["%d.%m.%y", "%d.%m.%Y"]

def _parse_date(v) -> str | None:
    """Parse DD.MM.YY or DD.MM.YYYY → YYYY-MM-DD."""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = _clean(v)
    if not s:
        return None
    for fmt in _BL_DATE_FMTS:
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


_PLZ_STADT = re.compile(r"^(\d{4,5})\s+(.+)$")

def _parse_address(raw: str) -> dict:
    """
    Parse a comma-separated address string into components.

    Examples:
      "Hans-Böckler-Straße 31, 65468 Trebur, , Deutschland"
        → address="Hans-Böckler-Straße 31", zip_code="65468", city="Trebur", country="DE"
      "In Kückhoven 11a,  Erkelenz, ,"
        → address="In Kückhoven 11a", city="Erkelenz"
    """
    parts = [p.strip() for p in raw.split(",")]
    result: dict = {}
    if not parts:
        return result

    if parts[0]:
        result["address"] = parts[0]

    for part in parts[1:]:
        part = part.strip()
        if not part:
            continue
        m = _PLZ_STADT.match(part)
        if m and "zip_code" not in result:
            result["zip_code"] = m.group(1)
            result["city"] = m.group(2).strip()
            continue
        country = transform_country(part)
        if country and "country" not in result:
            result["country"] = country
            continue
        if "city" not in result:
            result["city"] = part

    return result


def _split_guest_name(gast: str) -> tuple[str, str]:
    """
    Split "Vorname Nachname" into (first_name, last_name).
    First whitespace-delimited token → first_name; remainder → last_name.
    Single token (e.g. surname only) → ("", token).
    """
    s = gast.strip()
    if not s:
        return "", ""
    parts = s.split(" ", 1)
    if len(parts) == 2:
        return parts[0].strip(), parts[1].strip()
    return "", parts[0]


_STATUS_EXTRA = {
    "gebucht": "confirmed",
    "bestätigt": "confirmed",
}

def _map_status(raw: str) -> str:
    s = raw.strip().lower()
    return _STATUS_MAP.get(s) or _STATUS_EXTRA.get(s) or ""


# ── Sheet loading ─────────────────────────────────────────────────────────────

def _load_rows(file_or_path) -> list[dict]:
    _seek(file_or_path)
    wb = openpyxl.load_workbook(file_or_path, read_only=True, data_only=True)
    ws = wb.active
    raw = list(ws.iter_rows(values_only=True))
    if not raw:
        return []
    headers = [str(h) if h is not None else "" for h in raw[0]]
    rows = []
    for row in raw[1:]:
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        rows.append(dict(zip(headers, row)))
    return rows


# ── Reservation reader ────────────────────────────────────────────────────────

def _read_reservations(rows: list[dict]) -> list[dict]:
    result = []
    for row in rows:
        gast = _clean(row.get("Gast"))
        if not gast:
            continue

        if is_company_name(gast):
            first, last = "", gast
        else:
            first, last = _split_guest_name(gast)

        sys: dict = {}
        if not first:
            sys["first_name"] = ""
        if not last:
            sys["last_name"] = ""

        r: dict = {
            "first_name": first or "-",
            "last_name":  last  or "-",
            "Check In":   _parse_date(row.get("Anreise")),
            "Check Out":  _parse_date(row.get("Abreise")),
            "Zimmertyp":  _clean(row.get("Unterkunft")),
        }

        email = _clean(row.get("E-Mail"))
        if email:
            r["email"] = email

        phone = _clean(row.get("Telefon"))
        if phone:
            r["phone"] = phone

        preis = row.get("Preis")
        if preis is not None and str(preis).strip():
            r["Summe"] = str(preis)

        status = _map_status(_clean(row.get("Status", "")))
        if status:
            r["Status"] = status

        if sys:
            r["_system_changes"] = sys

        result.append({k: v for k, v in r.items() if v is not None and v != ""})
    return result


# ── Guest reader ──────────────────────────────────────────────────────────────

def _read_guests(rows: list[dict]) -> list[dict]:
    seen: dict[tuple, bool] = {}
    result = []
    for row in rows:
        gast = _clean(row.get("Gast"))
        if not gast or is_company_name(gast):
            continue

        first, last = _split_guest_name(gast)
        if not last:
            continue

        email = _clean(row.get("E-Mail"))
        key = (first.lower(), last.lower(), email.lower())
        if key in seen:
            continue
        seen[key] = True

        sys: dict = {}
        if not first:
            sys["first_name"] = ""

        g: dict = {
            "first_name": first or "-",
            "last_name":  last,
        }
        if email:
            g["email"] = email

        phone = _clean(row.get("Telefon"))
        if phone:
            g["phone"] = phone

        addr_raw = _clean(row.get("Adresse", ""))
        if addr_raw:
            g.update(_parse_address(addr_raw))

        if sys:
            g["_system_changes"] = sys

        result.append(g)
    return result


# ── Entry point ───────────────────────────────────────────────────────────────

def read_bookinglist(file_or_path, entity_type: str) -> list[dict]:
    """
    Read a BookingList XLSX and return rows pre-mapped to HotelFriend import format.

    file_or_path: path string, Path, or file-like object (BytesIO)
    entity_type:  "reservation" | "guest"
    """
    rows = _load_rows(file_or_path)

    if entity_type == "reservation":
        return _read_reservations(rows)
    if entity_type == "guest":
        return _read_guests(rows)
    return []
